"""Exportiert einen Auszug der Rezeptionszeugnisse als kompakte Liste (DOCX)."""

from __future__ import annotations

import argparse
import logging
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable, Iterator

from lxml import etree

TEI = "{http://www.tei-c.org/ns/1.0}"
NS = {"t": TEI[1:-1]}

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_INPUT = ROOT / "data" / "lists" / "sappho-rez_alle.xml"
OUTPUT = ROOT / "data" / "auszug_rezeptionszeugnisse.docx"

FIELDS = ("year", "author", "title", "container", "place", "genre")
DEFAULT_FIELDS = FIELDS
CENTURY_UNKNOWN = "Ohne Datierung"

log = logging.getLogger("reception-list")


@dataclass(frozen=True)
class Entry:
    id: str
    title: str
    authors: tuple[str, ...]
    genre: str
    year_label: str
    sort_year: int | None
    container: str
    place: str

    @property
    def author_label(self) -> str:
        return " / ".join(self.authors) if self.authors else "o. A."

    @property
    def century(self) -> int | None:
        return (self.sort_year - 1) // 100 + 1 if self.sort_year else None


def _text(el: etree._Element | None) -> str:
    return re.sub(r"\s+", " ", "".join(el.itertext())).strip() if el is not None else ""


def _year(date: etree._Element) -> int | None:
    for attr in ("when", "notBefore", "notAfter"):
        match = re.match(r"-?\d{3,4}", date.get(attr, ""))
        if match:
            return int(match.group())
    match = re.search(r"\d{4}", _text(date))
    return int(match.group()) if match else None


DATE_ABBREVIATIONS = {"published": "P", "created": "E"}

LEGEND = (
    "(E) = Entstehungsdatum, (P) = Publikationsdatum. "
    "Die Texte sind nach dem frühesten bekannten Jahr gelistet. "
    "Die Angaben sind nicht bei jedem Eintrag vollständig. "
    "n. n. = nicht nach, o. T. = ohne Titel, o. A. = ohne Autor_innenangabe. "
    "Das vollständige, laufend aktualisierte Verzeichnis ist auf Sappho Digital abrufbar "
    "(https://sappho-digital.com)."
)


def _pick_date(bibl: etree._Element) -> tuple[str, int | None]:
    """Beschriftet alle Datierungen mit ihrer Art, früheste zuerst; das früheste Jahr dient zum Sortieren."""
    dated = []
    for date in bibl.findall("t:date", NS):
        kind = date.get("type")
        year = _year(date)
        if kind in DATE_ABBREVIATIONS and year is not None:
            dated.append((year, kind != "published", _text(date), DATE_ABBREVIATIONS[kind]))
    if not dated:
        return "", None
    labels: dict[str, list[str]] = {}
    for _, _, text, abbreviation in sorted(dated):
        labels.setdefault(text, []).append(abbreviation)
    label = ", ".join(f"{text} ({'/'.join(kinds)})" for text, kinds in labels.items())
    return label, min(dated)[0]


def parse_entries(path: Path) -> list[Entry]:
    # xml:ids von Sammelwerken kommen mehrfach vor, daher tolerantes Parsen
    parser = etree.XMLParser(recover=True, remove_blank_text=True)
    tree = etree.parse(str(path), parser)
    entries = []
    for bibl in tree.iterfind(".//t:listBibl/t:bibl", NS):
        year_label, sort_year = _pick_date(bibl)
        entries.append(
            Entry(
                id=bibl.get("{http://www.w3.org/XML/1998/namespace}id", ""),
                title=_text(bibl.find("t:title", NS)),
                authors=tuple(a for a in (_text(e) for e in bibl.findall("t:author", NS)) if a),
                genre=_text(bibl.find("t:note[@type='genre']", NS)),
                year_label=year_label,
                sort_year=sort_year,
                container=_text(bibl.find("t:bibl/t:title", NS)),
                place=_text(bibl.find("t:pubPlace", NS)),
            )
        )
    log.info("%d Einträge gelesen aus %s", len(entries), path.name)
    return entries


def _norm(value: str) -> str:
    return value.casefold().strip()


def select(
    entries: Iterable[Entry],
    genres: list[str],
    authors: list[str],
    year_from: int | None,
    year_to: int | None,
    sample: int | None,
) -> list[Entry]:
    wanted_genres = {_norm(g) for g in genres}
    wanted_authors = [_norm(a) for a in authors]

    def keep(e: Entry) -> bool:
        if wanted_genres and _norm(e.genre.rstrip("?")) not in wanted_genres:
            return False
        if wanted_authors and not any(a in _norm(e.author_label) for a in wanted_authors):
            return False
        if year_from is not None or year_to is not None:
            if e.sort_year is None:
                return False
            if year_from is not None and e.sort_year < year_from:
                return False
            if year_to is not None and e.sort_year > year_to:
                return False
        return True

    chosen = sorted(
        (e for e in entries if keep(e)),
        key=lambda e: (e.sort_year is None, e.sort_year or 0, _norm(e.author_label), _norm(e.title)),
    )
    if sample and sample < len(chosen):
        chosen = _evenly_spaced(chosen, sample)
    return chosen


def _evenly_spaced(items: list[Entry], n: int) -> list[Entry]:
    """Deterministische Stichprobe, die die chronologische Verteilung beibehält."""
    step = len(items) / n
    return [items[int(i * step + step / 2)] for i in range(n)]


def _sentence(text: str) -> str:
    return text if text.endswith((".", "!", "?")) else f"{text}."


def entry_text(e: Entry, fields: tuple[str, ...]) -> tuple[str, str]:
    """Baut die Eintragszeile aus den gewählten Feldern: (fette Datierung, Rest)."""
    year = e.year_label if "year" in fields else ""
    parts: list[str] = ["\u2002"] if year else []
    add = parts.append

    if "author" in fields:
        add(e.author_label + (": " if "title" in fields else ". "))
    if "title" in fields:
        add(_sentence(e.title) + " " if e.title else "o. T. ")
    if "container" in fields and e.container:
        add(f"In: {_sentence(e.container)} ")
    if "place" in fields and e.place:
        add(_sentence(e.place) + " ")
    if "genre" in fields and e.genre:
        add(f"[{e.genre}]")
    return year, "".join(parts).rstrip()


def group_key(group_by: str) -> Callable[[Entry], str]:
    if group_by == "genre":
        return lambda e: e.genre or "Ohne Gattung"
    if group_by == "century":
        return lambda e: f"{e.century}. Jahrhundert" if e.century else CENTURY_UNKNOWN
    return lambda e: ""


def grouped(entries: list[Entry], group_by: str) -> Iterator[tuple[str, list[Entry]]]:
    key = group_key(group_by)
    groups: dict[str, list[Entry]] = {}
    for e in entries:
        groups.setdefault(key(e), []).append(e)
    yield from groups.items()


# Formatvorlagen wie in der Dissertation (Werte in Twips bzw. halben Punkten), damit die Liste beim
# Einfügen 1:1 übernommen wird; gleichnamige Vorlagen im Zieldokument greifen dann identisch.
STYLES_XML = """
<w:style xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    w:type="paragraph" w:customStyle="1" w:styleId="Rezeptionsliste">
  <w:name w:val="Rezeptionsliste"/><w:basedOn w:val="Normal"/><w:qFormat/>
  <w:pPr><w:keepLines/><w:suppressAutoHyphens/>
    <w:spacing w:before="0" w:after="{after}" w:line="250" w:lineRule="atLeast"/>
    <w:ind w:left="454" w:hanging="454"/><w:jc w:val="left"/></w:pPr>
  <w:rPr><w:sz w:val="{size}"/><w:szCs w:val="{size}"/></w:rPr>
</w:style>
<w:style xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    w:type="paragraph" w:customStyle="1" w:styleId="RezeptionslisteHinweis">
  <w:name w:val="Rezeptionsliste Hinweis"/><w:basedOn w:val="Normal"/><w:qFormat/>
  <w:pPr><w:keepLines/><w:suppressAutoHyphens/>
    <w:spacing w:before="0" w:after="240" w:line="270" w:lineRule="atLeast"/>
    <w:ind w:firstLine="0"/><w:jc w:val="left"/></w:pPr>
  <w:rPr><w:i/><w:iCs/><w:sz w:val="{size}"/><w:szCs w:val="{size}"/></w:rPr>
</w:style>
"""


def _setup_document(font: str, size: float, spacing: float):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Twips

    doc = Document()

    # Seitenformat der Dissertation (A4, Satzspiegel)
    section = doc.sections[0]
    section.page_width, section.page_height = Twips(11906), Twips(16838)
    section.top_margin, section.bottom_margin = Twips(1644), Twips(1814)
    section.left_margin = section.right_margin = Twips(2211)
    section.header_distance, section.footer_distance = Twips(850), Twips(907)

    # Standard wie in der Dissertation
    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    rfonts = normal.element.rPr.rFonts
    for attr in ("w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font)
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "de-DE")
    normal.element.rPr.append(lang)
    fmt = normal.paragraph_format
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    fmt.line_spacing = Twips(330)
    fmt.first_line_indent = Twips(340)
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    half_points = round(size * 2)
    styles = parse_xml(f"<w:root xmlns:w='http://schemas.openxmlformats.org/wordprocessingml/2006/main'>"
                       f"{STYLES_XML.format(size=half_points, after=round(spacing * 20))}</w:root>")
    for style in list(styles):
        doc.styles.element.append(style)
    return doc


def write_docx(
    path: Path,
    entries: list[Entry],
    fields: tuple[str, ...],
    group_by: str,
    font: str,
    size: float,
    spacing: float,
) -> None:
    doc = _setup_document(font, size, spacing)
    doc.add_paragraph(LEGEND, style="Rezeptionsliste Hinweis")

    for heading, group in grouped(entries, group_by):
        if heading:
            p = doc.add_paragraph(style="Rezeptionsliste")
            p.paragraph_format.left_indent = 0
            p.paragraph_format.first_line_indent = 0
            p.paragraph_format.keep_with_next = True
            p.add_run(heading).bold = True
        for e in group:
            year, rest = entry_text(e, fields)
            p = doc.add_paragraph(style="Rezeptionsliste")
            if year:
                p.add_run(year).bold = True
            p.add_run(rest)
    doc.save(path)


def _csv_list(value: str) -> list[str]:
    return [v.strip() for v in value.split(",") if v.strip()]


def _fields(value: str) -> tuple[str, ...]:
    chosen = tuple(_csv_list(value))
    unknown = set(chosen) - set(FIELDS)
    if unknown:
        raise argparse.ArgumentTypeError(
            f"Unbekannte Felder: {', '.join(sorted(unknown))}. Erlaubt: {', '.join(FIELDS)}"
        )
    return chosen


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("-i", "--input", type=Path, default=DEFAULT_INPUT, help="TEI-Liste (Standard: sappho-rez_alle.xml)")
    p.add_argument("--fields", type=_fields, default=DEFAULT_FIELDS, metavar="F1,F2,…", help=f"Felder in Ausgabereihenfolge-Logik; wählbar: {', '.join(FIELDS)}")
    p.add_argument("--genre", type=_csv_list, default=[], metavar="G1,G2", help="Nur diese Gattungen, z. B. Lyrik,Drama")
    p.add_argument("--author", type=_csv_list, default=[], metavar="A1,A2", help="Nur Autor:innen, deren Name den Text enthält")
    p.add_argument("--from", dest="year_from", type=int, help="frühestes Jahr")
    p.add_argument("--to", dest="year_to", type=int, help="spätestes Jahr")
    p.add_argument("--sample", type=int, metavar="N", help="N chronologisch gleichmäßig verteilte Einträge")
    p.add_argument("--group-by", choices=("none", "genre", "century"), default="none", help="Zwischenüberschriften")
    p.add_argument("--font", default="EB Garamond")
    p.add_argument("--size", type=float, default=10.0, help="Schriftgröße in pt")
    p.add_argument("--spacing", type=float, default=3.5, help="Abstand nach Eintrag in pt")
    return p


def main(argv: list[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    args = build_parser().parse_args(argv)

    if not args.input.is_file():
        log.error("Eingabedatei nicht gefunden: %s", args.input)
        return 1

    entries = select(
        parse_entries(args.input),
        genres=args.genre,
        authors=args.author,
        year_from=args.year_from,
        year_to=args.year_to,
        sample=args.sample,
    )
    if not entries:
        log.error("Keine Einträge entsprechen den Filtern.")
        return 1

    write_docx(OUTPUT, entries, args.fields, args.group_by, args.font, args.size, args.spacing)

    log.info("%d Einträge geschrieben nach %s", len(entries), OUTPUT)
    return 0


if __name__ == "__main__":
    sys.exit(main())
